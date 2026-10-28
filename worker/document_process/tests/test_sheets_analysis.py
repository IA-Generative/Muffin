import datetime as dt
import json
import tracemalloc
import zipfile
from pathlib import Path

import pytest
import xlwt
from openpyxl import Workbook

from app.tabular.sheets.analyzer import analyze_file
from app.tabular.sheets.limits import Limits
from app.tabular.sheets.readers import UnreadableFile, UnsupportedFormat

ODS_HEAD = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet>'
)
ODS_TAIL = "</office:spreadsheet></office:body></office:document-content>"


def write_ods(path: Path, tables: dict[str, str]) -> Path:
    body = "".join(f'<table:table table:name="{name}">{rows}</table:table>' for name, rows in tables.items())
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        archive.writestr("content.xml", ODS_HEAD + body + ODS_TAIL)
    return path


def ods_text(value: str) -> str:
    return f'<table:table-cell office:value-type="string"><text:p>{value}</text:p></table:table-cell>'


def ods_float(value: float) -> str:
    cell = f'<table:table-cell office:value-type="float" office:value="{value}">'
    return f"{cell}<text:p>{value}</text:p></table:table-cell>"


def ods_row(*cells: str) -> str:
    return f"<table:table-row>{''.join(cells)}</table:table-row>"


def column(sheet: dict, name: str) -> dict:
    return next(c for c in sheet["columns"] if c["name"] == name)


# --- CSV ---


def test_csv_semicolon_french_formats(tmp_path) -> None:
    path = tmp_path / "ventes.csv"
    path.write_text(
        "produit;prix;date;actif;code\nStylo;1,50;01/02/2024;vrai;007\nCahier;3,2;15/03/2024;faux;012\n"
        "Stylo;;20/04/2024;vrai;3\n",
        encoding="utf-8",
    )

    sheet = analyze_file(path, "ventes.csv")["sheets"][0]

    assert (sheet["rows"], sheet["header_detected"]) == (3, True)
    assert {c["name"]: c["type"] for c in sheet["columns"]} == {
        "produit": "text",
        "prix": "number",
        "date": "date",
        "actif": "boolean",
        "code": "mixed",  # '007' and '012' stay text (leading zeros), '3' is a number
    }
    price = column(sheet, "prix")
    assert (price["non_null"], price["null"]) == (2, 1)
    assert price["numeric"]["min"] == 1.5 and price["numeric"]["max"] == 3.2
    assert column(sheet, "date")["dates"] == {"min": "2024-02-01", "max": "2024-04-20"}
    assert column(sheet, "produit")["top_values"][0] == {"value": "Stylo", "count": 2}


def test_csv_comma_delimited_keeps_dot_decimals_and_latin1(tmp_path) -> None:
    path = tmp_path / "x.csv"
    path.write_bytes("ville,montant\nLyon,10.5\nMarseille,20\nÉvry,30.25\n".encode("cp1252"))

    sheet = analyze_file(path, "x.csv")["sheets"][0]

    assert [c["name"] for c in sheet["columns"]] == ["ville", "montant"]
    assert column(sheet, "ville")["top_values"][-1]["value"] == "Évry"
    assert column(sheet, "montant")["numeric"]["sum"] == 60.75


def test_csv_without_header_gets_generated_names_and_ragged_rows(tmp_path) -> None:
    path = tmp_path / "raw.csv"
    path.write_text("1,2,3\n4,5\n6,7,8,9\n", encoding="utf-8")

    sheet = analyze_file(path, "raw.csv")["sheets"][0]

    assert sheet["header_detected"] is False
    assert [c["name"] for c in sheet["columns"]] == ["col_1", "col_2", "col_3", "col_4"]
    assert (sheet["rows"], column(sheet, "col_3")["null"], column(sheet, "col_4")["non_null"]) == (3, 1, 1)


def test_duplicate_and_empty_headers_are_made_unique(tmp_path) -> None:
    path = tmp_path / "d.csv"
    path.write_text("a,a,,b\n1,2,3,4\n", encoding="utf-8")

    names = [c["name"] for c in analyze_file(path, "d.csv")["sheets"][0]["columns"]]

    assert names == ["a", "a_2", "col_3", "b"]


def test_empty_file_is_reported_without_crashing(tmp_path) -> None:
    path = tmp_path / "empty.csv"
    path.write_text("", encoding="utf-8")

    result = analyze_file(path, "empty.csv")

    assert result["summary"]["rows"] == 0 and "The file contains no data" in result["warnings"]


# --- XLSX ---


def test_xlsx_types_dates_errors_and_sheets(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventes"
    sheet.append(["mois", "montant", "date", "ok"])
    sheet.append(["jan", 100, dt.datetime(2024, 1, 5), True])
    sheet.append(["fev", 250.5, dt.datetime(2024, 2, 9), False])
    sheet.append(["mar", "#DIV/0!", dt.datetime(2024, 3, 1), True])
    hidden = workbook.create_sheet("Cache")
    hidden.sheet_state = "hidden"
    hidden.append(["x"])
    hidden.append([1])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    result = analyze_file(path, "book.xlsx")

    sheets = {s["name"]: s for s in result["sheets"]}
    assert result["summary"]["sheets"] == 2 and sheets["Cache"]["hidden"] is True
    amount = column(sheets["Ventes"], "montant")
    assert amount["type"] == "number" and amount["error_cells"] == 1
    assert amount["numeric"]["count"] == 2
    assert column(sheets["Ventes"], "date")["type"] == "datetime"
    assert column(sheets["Ventes"], "ok")["type"] == "boolean"
    json.dumps(result, allow_nan=False)  # the whole result must be strict JSON


def test_xlsx_with_a_bogus_huge_dimension_is_still_read_quickly(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["a", "b"])
    sheet.append([1, 2])
    sheet["A1"].value = "a"
    path = tmp_path / "dim.xlsx"
    workbook.save(path)
    # Rewrite the sheet XML with a declared dimension of the whole grid.
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b'<dimension ref="A1:B2"/>', b'<dimension ref="A1:XFD1048576"/>'
    )
    with zipfile.ZipFile(path, "w") as target:
        for name, data in parts.items():
            target.writestr(name, data)

    sheet_result = analyze_file(path, "dim.xlsx")["sheets"][0]

    assert sheet_result["rows"] == 1 and len(sheet_result["columns"]) == 2


# --- ODS ---


def test_ods_values_repeats_and_trailing_blanks(tmp_path) -> None:
    rows = (
        ods_row(ods_text("produit"), ods_text("prix"))
        + ods_row(ods_text("Stylo"), ods_float(1.5))
        + '<table:table-row table:number-rows-repeated="3">'
        + ods_text("Cahier")
        + ods_float(3)
        + "</table:table-row>"
        # the usual padding: a huge run of empty cells and rows that must not be expanded
        + '<table:table-row><table:table-cell table:number-columns-repeated="16384"/></table:table-row>'
        + '<table:table-row table:number-rows-repeated="1048570"><table:table-cell '
        'table:number-columns-repeated="16384"/></table:table-row>'
    )
    path = write_ods(tmp_path / "t.ods", {"Feuille1": rows, "Feuille2": ods_row(ods_text("a")) + ods_row(ods_float(1))})

    result = analyze_file(path, "t.ods")

    first, second = result["sheets"]
    assert (first["name"], first["rows"], len(first["columns"])) == ("Feuille1", 4, 2)
    price = column(first, "prix")
    assert price["numeric"]["count"] == 4 and price["numeric"]["sum"] == 1.5 + 9
    assert column(first, "produit")["top_values"][0] == {"value": "Cahier", "count": 3}
    assert (second["name"], second["rows"]) == ("Feuille2", 1)


def test_ods_dates_booleans_and_covered_cells(tmp_path) -> None:
    date = '<table:table-cell office:value-type="date" office:date-value="2024-05-17"/>'
    boolean = '<table:table-cell office:value-type="boolean" office:boolean-value="true"/>'
    rows = (
        ods_row(ods_text("jour"), ods_text("ok"))
        + ods_row(date, boolean)
        + ods_row("<table:covered-table-cell/>", boolean)
    )
    path = write_ods(tmp_path / "d.ods", {"S": rows})

    sheet = analyze_file(path, "d.ods")["sheets"][0]

    assert column(sheet, "jour")["dates"] == {"min": "2024-05-17", "max": "2024-05-17"}
    assert column(sheet, "ok")["type"] == "boolean"


def test_stopping_after_a_sheet_does_not_break_the_next_ones(tmp_path) -> None:
    big = "".join(ods_row(ods_float(i)) for i in range(50))
    path = write_ods(
        tmp_path / "s.ods", {"A": ods_row(ods_text("n")) + big, "B": ods_row(ods_text("m")) + ods_row(ods_float(1))}
    )

    result = analyze_file(path, "s.ods", Limits(max_rows_per_sheet=10))

    first, second = result["sheets"]
    assert first["truncated"] is True and first["rows"] == 10
    assert (second["name"], second["rows"]) == ("B", 1)


# --- XLS ---


def test_xls_numbers_dates_and_text(tmp_path) -> None:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Legacy")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for col, label in enumerate(["nom", "age", "naissance"]):
        sheet.write(0, col, label)
    sheet.write(1, 0, "Alice")
    sheet.write(1, 1, 30)
    sheet.write(1, 2, dt.datetime(1994, 7, 14), date_style)
    sheet.write(2, 0, "Bob")
    sheet.write(2, 1, 41)
    sheet.write(2, 2, dt.datetime(1983, 1, 2), date_style)
    path = tmp_path / "old.xls"
    workbook.save(str(path))

    sheet_result = analyze_file(path, "old.xls")["sheets"][0]

    assert sheet_result["name"] == "Legacy" and sheet_result["rows"] == 2
    assert column(sheet_result, "age")["type"] == "integer"
    assert column(sheet_result, "naissance")["dates"]["min"].startswith("1983-01-02")


# --- statistics ---


def test_numeric_statistics_are_exact_for_small_columns(tmp_path) -> None:
    path = tmp_path / "n.csv"
    path.write_text("v\n" + "\n".join(str(i) for i in range(1, 101)) + "\n", encoding="utf-8")

    stats = column(analyze_file(path, "n.csv")["sheets"][0], "v")["numeric"]

    assert (stats["min"], stats["max"], stats["sum"], stats["mean"]) == (1, 100, 5050, 50.5)
    assert stats["median"] == 50.5 and stats["quantiles_exact"] is True
    assert stats["std"] == pytest.approx(29.011491975882016)


def test_quantiles_become_estimates_past_the_reservoir_and_distinct_is_capped(tmp_path) -> None:
    path = tmp_path / "big.csv"
    path.write_text("id,v\n" + "\n".join(f"u{i},{i}" for i in range(5000)) + "\n", encoding="utf-8")
    limits = Limits(reservoir_size=500, max_distinct=100)

    sheet = analyze_file(path, "big.csv", limits)["sheets"][0]

    numeric = column(sheet, "v")["numeric"]
    assert numeric["quantiles_exact"] is False
    assert numeric["median"] == pytest.approx(2500, rel=0.15)  # uniform sample estimate
    assert (numeric["min"], numeric["max"], numeric["count"]) == (0, 4999, 5000)  # these stay exact
    ids = column(sheet, "id")
    assert ids["distinct"] == 100 and ids["distinct_exact"] is False


def test_row_and_column_limits_flag_truncation(tmp_path) -> None:
    path = tmp_path / "w.csv"
    path.write_text("a,b,c,d\n" + "\n".join("1,2,3,4" for _ in range(20)) + "\n", encoding="utf-8")

    sheet = analyze_file(path, "w.csv", Limits(max_rows_per_sheet=5, max_columns=2))["sheets"][0]

    assert (sheet["rows"], sheet["truncated"], len(sheet["columns"])) == (5, True, 2)
    assert len(sheet["warnings"]) == 2


def test_output_size_does_not_depend_on_the_number_of_rows(tmp_path) -> None:
    sizes = []
    for rows in (100, 20_000):
        path = tmp_path / f"{rows}.csv"
        path.write_text("a,b\n" + "\n".join(f"x{i},{i}" for i in range(rows)) + "\n", encoding="utf-8")
        sizes.append(len(json.dumps(analyze_file(path, path.name))))

    assert sizes[1] < sizes[0] * 4


# --- errors ---


def test_unsupported_and_corrupt_files(tmp_path) -> None:
    exe = tmp_path / "a.exe"
    exe.write_bytes(b"MZ")
    fake = tmp_path / "fake.xlsx"
    fake.write_text("not a zip", encoding="utf-8")
    broken_ods = tmp_path / "broken.ods"
    broken_ods.write_bytes(b"PK\x03\x04garbage")
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * 5_000_000)

    with pytest.raises(UnsupportedFormat):
        analyze_file(exe, "a.exe")
    with pytest.raises(UnreadableFile, match="not a zip"):
        analyze_file(fake, "fake.xlsx")
    with pytest.raises(UnreadableFile):
        analyze_file(broken_ods, "broken.ods")
    with pytest.raises(UnreadableFile, match="expands"):
        analyze_file(bomb, "bomb.xlsx", Limits(max_uncompressed_bytes=1_000_000))


# --- memory: peak allocation must not grow with the number of rows ---


def peak_mb(path: Path, name: str) -> float:
    tracemalloc.start()
    try:
        analyze_file(path, name)
        return tracemalloc.get_traced_memory()[1] / 1e6
    finally:
        tracemalloc.stop()


def test_csv_memory_is_flat(tmp_path) -> None:
    def make(rows: int) -> Path:
        path = tmp_path / f"m{rows}.csv"
        with path.open("w") as handle:
            handle.write("id,nom,montant,date\n")
            handle.writelines(f"{i},client{i % 5000},{i * 1.5},2024-01-{i % 28 + 1:02d}\n" for i in range(rows))
        return path

    small, large = peak_mb(make(5_000), "s.csv"), peak_mb(make(300_000), "l.csv")

    assert large < 25  # 300k rows; holding them would take hundreds of MB
    assert large < small * 3


def test_xlsx_memory_is_flat(tmp_path) -> None:
    def make(rows: int) -> Path:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("data")
        sheet.append(["id", "nom", "montant"])
        for i in range(rows):
            sheet.append([i, f"client{i % 3000}", i * 1.5])
        path = tmp_path / f"m{rows}.xlsx"
        workbook.save(path)
        return path

    small, large = peak_mb(make(2_000), "s.xlsx"), peak_mb(make(100_000), "l.xlsx")

    # openpyxl's incremental parser leaves some XML bookkeeping per row: 78-83 bytes measured on its own, and up to
    # ~150 under pytest on Python 3.13 (the interpreter's allocator accounting differs). Loading the cells (the
    # non-streaming mode) would cost several hundred bytes per *cell*, three cells a row here - so 250 still tells a
    # stream from a load.
    assert large < 25
    assert (large - small) * 1e6 / 98_000 < 250


def test_ods_memory_is_flat(tmp_path) -> None:
    def make(rows: int) -> Path:
        body = ods_row(ods_text("id"), ods_text("nom")) + "".join(
            ods_row(ods_float(i), ods_text(f"c{i % 2000}")) for i in range(rows)
        )
        return write_ods(tmp_path / f"m{rows}.ods", {"data": body})

    small, large = peak_mb(make(2_000), "s.ods"), peak_mb(make(100_000), "l.ods")

    assert large < 40
    assert large < small * 4

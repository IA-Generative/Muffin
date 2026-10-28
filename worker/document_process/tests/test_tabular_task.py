"""process_tabular_document, from the file to the chunking, on real files: only the backend, S3 and the model are
played."""

import io
from unittest.mock import MagicMock

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from app.tabular.sheets import compat
from app.tasks import _tabular_steps as steps
from app.tasks import tabular


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    workbook = Workbook(write_only=True)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def world(monkeypatch):
    """The backend, the storage and the model, played. `world.put(name, bytes)` is the file in S3."""
    shared = steps._shared

    class World:
        def __init__(self):
            self.files: dict[str, bytes] = {}
            self.uploaded: dict[str, bytes] = {}
            self.document = {"id": "doc-1", "name": "f.xlsx", "storage_key": "docs/f", "mime_type": ""}
            self.backend = MagicMock()
            self.storage = MagicMock()

        def put(self, name: str, data: bytes, mime: str = "") -> None:
            self.document.update({"name": name, "mime_type": mime})
            self.files["docs/f"] = data

    w = World()
    w.backend.get_document.side_effect = lambda document_id: w.document
    w.backend.get_collection_settings.return_value = {
        "embedding_model": "m",
        "instructions": {"summary": "", "qa": ""},
    }
    w.storage.head_object.side_effect = lambda key: {"ContentLength": len(w.files[key])}
    w.storage.download_fileobj.side_effect = lambda key, handle: handle.write(w.files[key])
    w.storage.upload_fileobj.side_effect = lambda key, handle, content_type="": w.uploaded.__setitem__(
        key, handle.read()
    )
    monkeypatch.setattr(shared, "backend_client", w.backend)
    monkeypatch.setattr(shared, "storage", w.storage)
    monkeypatch.setattr(shared, "_model_for", MagicMock(return_value="gpt"))
    monkeypatch.setattr(shared, "_windows", MagicMock(return_value={"qa_questions_per_window": 2}))
    monkeypatch.setattr(
        shared,
        "_chat",
        MagicMock(
            side_effect=lambda model, system, user: (
                '[{"question": "q", "answer": "a"}]' if "JSON" in system else "résumé"
            )
        ),
    )
    monkeypatch.setattr(shared, "_spawn", MagicMock())
    monkeypatch.setattr(shared, "_extraction_done", MagicMock())
    monkeypatch.setattr(shared, "_fail", MagicMock())
    w.spawn = shared._spawn
    return w


def _profile_sent(world) -> dict:
    return world.backend.set_tabular_profile.call_args.args[1]


def test_a_workbook_gives_a_profile_a_typed_copy_and_a_page_per_sheet(world):
    world.put(
        "ventes.xlsx",
        _xlsx(
            {
                "Ventes": [["region", "montant"], ["Nord", 10], ["Sud", 20], ["Nord", 30]],
                "Clients": [["id", "nom"], [1, "Alice"], [2, "Bob"]],
            }
        ),
    )

    tabular.process_tabular_document("doc-1", "col-1", "xlsx")

    profile = _profile_sent(world)
    assert (profile["document_id"], profile["format"], profile["row_count"], profile["column_count"]) == (
        "doc-1",
        "xlsx",
        3,
        2,
    )
    assert [column["name"] for column in profile["columns"]] == ["region", "montant"]
    assert profile["analysis_version"] == compat.ANALYSIS_VERSION
    sheets = profile["sheets"]
    assert [(s["index"], s["name"], s["rows"]) for s in sheets] == [(0, "Ventes", 3), (1, "Clients", 2)]
    assert sheets[0]["parquet_key"] == "documents/col-1/doc-1/tables/0.parquet"
    assert sheets[1]["table"]["columns"] == [
        {"name": "id", "sql_type": "BIGINT"},
        {"name": "nom", "sql_type": "VARCHAR"},
    ]
    assert set(world.uploaded) == {"documents/col-1/doc-1/tables/0.parquet", "documents/col-1/doc-1/tables/1.parquet"}
    assert pq.read_table(io.BytesIO(world.uploaded["documents/col-1/doc-1/tables/1.parquet"])).num_rows == 2
    pages = {c.kwargs["page_number"]: c.kwargs["content"] for c in world.backend.add_page.call_args_list}
    assert set(pages) == {1, 2} and pages[1].startswith("# Ventes\n")
    world.backend.prune_pages.assert_called_once_with("doc-1", keep=2)


def test_the_summary_and_the_questions_are_made_from_the_profile_and_the_chunking_is_skipped(world):
    world.put("ventes.xlsx", _xlsx({"Ventes": [["region", "montant"], ["Nord", 10], ["Sud", 20]]}))

    tabular.process_tabular_document("doc-1", "col-1", "xlsx")

    world.backend.set_document_summary.assert_called_once()
    assert world.backend.set_document_summary.call_args.args[:2] == ("doc-1", "résumé")
    world.backend.replace_generated_qa.assert_called_once()
    spawn = world.spawn.call_args
    assert spawn.args[2] == "app.tasks.chunk_document"
    assert spawn.args[1] == ["doc-1", "col-1", True, True, True, True]  # nothing left for chunking to do
    world.backend.update_status.assert_any_call("doc-1", status="indexing", progress=50)


def test_a_workbook_of_several_sheets_says_so_to_the_model(world):
    world.put("f.xlsx", _xlsx({"Ventes": [["a"], [1], [2]], "Clients": [["b"], [1]]}))

    tabular.process_tabular_document("doc-1", "col-1", "xlsx")

    prompts = [call.args[2] for call in steps._shared._chat.call_args_list]
    assert prompts and all("Le classeur contient 2 feuilles" in prompt for prompt in prompts)
    assert all("Clients (1 lignes)" in prompt and "« Ventes »" in prompt for prompt in prompts)


@pytest.mark.parametrize(
    ("name", "content", "fmt"),
    [
        ("data.csv", b"region,montant\nNord,10\nSud,20\n", "csv"),
        ("data.tsv", b"region\tmontant\nNord\t10\nSud\t20\n", "tsv"),
        ("data.psv", b"region|montant\nNord|10\nSud|20\n", "psv"),
        ("data.jsonl", b'{"region": "Nord", "montant": 10}\n{"region": "Sud", "montant": 20}\n', "jsonl"),
        ("data.json", b'[{"region": "Nord", "montant": 10}, {"region": "Sud", "montant": 20}]', "json"),
    ],
)
def test_the_formats_muffin_always_took_still_give_a_profile_and_a_copy(world, name, content, fmt):
    world.put(name, content)

    tabular.process_tabular_document("doc-1", "col-1", fmt)

    profile = _profile_sent(world)
    assert (profile["row_count"], profile["format"]) == (2, fmt)
    assert [column["name"] for column in profile["columns"]] == ["region", "montant"]
    assert [column["type"] for column in profile["columns"]] == ["VARCHAR", "BIGINT"]
    assert profile["sheets"][0]["name"] == "data" and len(world.uploaded) == 1


def test_a_file_known_only_by_its_content_is_read_as_the_format_it_was_recognized_as(world):
    """No extension, or a wrong one: the readers go by the extension, so the step gives them the right one."""
    world.put("export", b"region,montant\nNord,10\nSud,20\n")

    tabular.process_tabular_document("doc-1", "col-1", "csv")

    assert _profile_sent(world)["row_count"] == 2


def test_a_file_with_no_data_ends_the_document_in_error_with_the_reason(world):
    world.put("empty.csv", b"")

    with pytest.raises(ValueError, match="no data"):
        tabular.process_tabular_document("doc-1", "col-1", "csv")

    world.backend.update_status.assert_any_call("doc-1", status="error")
    steps._shared._fail.assert_called_once()
    assert "no data" in str(steps._shared._fail.call_args.args[2])
    world.backend.set_tabular_profile.assert_not_called()
    world.spawn.assert_not_called()


def test_a_file_over_the_limit_is_refused_before_it_is_downloaded(world, monkeypatch):
    monkeypatch.setenv("TABULAR_MAX_FILE_BYTES", "10")
    world.put("big.csv", b"region,montant\nNord,10\n")

    with pytest.raises(ValueError, match="File too large"):
        tabular.process_tabular_document("doc-1", "col-1", "csv")

    world.storage.download_fileobj.assert_not_called()
    world.backend.update_status.assert_any_call("doc-1", status="error")


def test_a_corrupt_workbook_is_refused_with_a_message_and_nothing_is_written(world):
    world.put("broken.xlsx", b"this is not a zip archive at all")

    with pytest.raises(ValueError, match="not a valid .xlsx"):
        tabular.process_tabular_document("doc-1", "col-1", "xlsx")

    assert not world.uploaded
    world.backend.update_status.assert_any_call("doc-1", status="error")


def test_a_document_with_an_unsupported_mime_type_is_refused(world):
    world.put("data.csv", b"a,b\n1,2\n", mime="application/pdf")
    with pytest.raises(ValueError, match="Unsupported MIME"):
        tabular.process_tabular_document("doc-1", "col-1", "csv")


def test_a_second_extraction_replaces_the_copies_and_the_pages_of_the_first(world):
    world.put("f.xlsx", _xlsx({"A": [["x"], [1]], "B": [["y"], [2]], "C": [["z"], [3]]}))
    tabular.process_tabular_document("doc-1", "col-1", "xlsx")
    world.uploaded.clear()
    world.backend.reset_mock()
    world.put("f.xlsx", _xlsx({"A": [["x"], [1]]}))

    tabular.process_tabular_document("doc-1", "col-1", "xlsx")

    world.storage.delete_prefix.assert_called_with("documents/col-1/doc-1/tables/")
    assert set(world.uploaded) == {"documents/col-1/doc-1/tables/0.parquet"}
    world.backend.prune_pages.assert_called_once_with("doc-1", keep=1)


def test_an_xls_workbook_is_read_like_an_xlsx(world, tmp_path):
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Ventes")
    for row, values in enumerate([["region", "montant"], ["Nord", 10], ["Sud", 20]]):
        for column, value in enumerate(values):
            sheet.write(row, column, value)
    path = tmp_path / "old.xls"
    workbook.save(str(path))
    world.put("old.xls", path.read_bytes(), mime="application/vnd.ms-excel")

    tabular.process_tabular_document("doc-1", "col-1", "xls")

    profile = _profile_sent(world)
    assert (profile["format"], profile["row_count"]) == ("xls", 2)
    assert profile["sheets"][0]["name"] == "Ventes"


def test_an_ods_workbook_is_read_like_an_xlsx(world, tmp_path):
    from tests.test_sheets_analysis import ods_float, ods_row, ods_text, write_ods

    rows = ods_row(ods_text("region"), ods_text("montant")) + ods_row(ods_text("Nord"), ods_float(10))
    rows += ods_row(ods_text("Sud"), ods_float(20))
    path = write_ods(tmp_path / "open.ods", {"Ventes": rows})
    world.put("open.ods", path.read_bytes(), mime="application/vnd.oasis.opendocument.spreadsheet")

    tabular.process_tabular_document("doc-1", "col-1", "ods")

    profile = _profile_sent(world)
    assert (profile["format"], profile["row_count"]) == ("ods", 2)
    assert [column["name"] for column in profile["columns"]] == ["region", "montant"]

import datetime as dt
from pathlib import Path

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from app.tabular.sheets.analyzer import analyze_file
from app.tabular.sheets.limits import Limits
from app.tabular.sheets.tables import column_kind, table_key, write_tables


def tables_of(path: Path, name: str, tmp_path: Path, limits: Limits | None = None):
    result = analyze_file(path, name, limits)
    out = tmp_path / "tables"
    out.mkdir(exist_ok=True)
    described = write_tables(path, name, result, out, limits)
    return result, described, out


def read(out: Path, index: int = 0) -> list[dict]:
    return pq.read_table(out / f"{index}.parquet").to_pylist()


def test_the_key_is_next_to_analyse_json_under_the_files_version() -> None:
    assert table_key("p/col/file/versions/2/analyse.json", 3) == "p/col/file/versions/2/tables/3.parquet"


def test_csv_columns_get_their_native_types(tmp_path) -> None:
    path = tmp_path / "v.csv"
    path.write_text(
        "produit;prix;qte;date;actif\nStylo;1,50;10;01/02/2024;vrai\nCahier;3,2;4;15/03/2024;faux\n", encoding="utf-8"
    )

    _, described, out = tables_of(path, "v.csv", tmp_path)

    assert [(c["name"], c["sql_type"]) for c in described[0]["columns"]] == [
        ("produit", "VARCHAR"),
        ("prix", "DOUBLE"),
        ("qte", "BIGINT"),
        ("date", "DATE"),
        ("actif", "BOOLEAN"),
    ]
    assert read(out)[0] == {
        "produit": "Stylo",
        "prix": 1.5,
        "qte": 10,
        "date": dt.date(2024, 2, 1),
        "actif": True,
    }


def test_the_header_is_not_data_and_a_headerless_sheet_keeps_its_first_row(tmp_path) -> None:
    with_header = tmp_path / "h.csv"
    with_header.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    without = tmp_path / "n.csv"
    without.write_text("1,2\n3,4\n", encoding="utf-8")

    _, d1, o1 = tables_of(with_header, "h.csv", tmp_path)
    (tmp_path / "t2").mkdir()
    _, d2, o2 = tables_of(without, "n.csv", tmp_path / "t2")

    assert d1[0]["rows"] == 2 and [r["a"] for r in read(o1)] == [1, 3]
    assert d2[0]["rows"] == 2 and [c["name"] for c in d2[0]["columns"]] == ["col_1", "col_2"]
    assert [r["col_1"] for r in read(o2)] == [1, 3]


def test_a_mixed_column_is_stored_as_text_so_that_no_value_is_lost(tmp_path) -> None:
    path = tmp_path / "m.csv"
    path.write_text("id,valeur\n1,10\n2,dix\n3,12\n", encoding="utf-8")

    _, described, out = tables_of(path, "m.csv", tmp_path)

    assert described[0]["columns"][1]["sql_type"] == "VARCHAR"
    assert [r["valeur"] for r in read(out)] == ["10", "dix", "12"]


def test_error_cells_become_null_in_a_typed_column(tmp_path) -> None:
    path = tmp_path / "e.csv"
    path.write_text("id,ratio\n1,0.5\n2,#DIV/0!\n3,2\n", encoding="utf-8")

    _, described, out = tables_of(path, "e.csv", tmp_path)

    assert described[0]["columns"][1]["sql_type"] == "DOUBLE"
    assert [r["ratio"] for r in read(out)] == [0.5, None, 2.0]


def test_xlsx_sheets_each_get_a_table_and_empty_sheets_none(tmp_path) -> None:
    book = Workbook()
    ventes = book.active
    ventes.title = "Ventes"
    ventes.append(["region", "montant", "jour"])
    ventes.append(["Nord", 10.5, dt.datetime(2024, 1, 2, 8, 30)])
    ventes.append(["Sud", 7, dt.datetime(2024, 1, 3, 9, 0)])
    book.create_sheet("Vide")
    rh = book.create_sheet("RH")
    rh.append(["nom", "embauche"])
    rh.append(["Ann", dt.date(2020, 5, 1)])
    path = tmp_path / "w.xlsx"
    book.save(path)

    result, described, out = tables_of(path, "w.xlsx", tmp_path)

    assert sorted(described) == [0, 2]  # sheet 1 ("Vide") has nothing to query
    assert [c["sql_type"] for c in described[0]["columns"]] == ["VARCHAR", "DOUBLE", "TIMESTAMP"]
    assert read(out, 0)[0]["jour"] == dt.datetime(2024, 1, 2, 8, 30)
    assert read(out, 2) == [{"nom": "Ann", "embauche": dt.datetime(2020, 5, 1)}]  # Excel has no date-only type
    assert result["sheets"][2]["name"] == "RH"


def test_a_truncated_sheet_is_truncated_in_the_same_place(tmp_path) -> None:
    path = tmp_path / "big.csv"
    path.write_text("n\n" + "\n".join(str(i) for i in range(50)) + "\n", encoding="utf-8")

    result, described, out = tables_of(path, "big.csv", tmp_path, Limits(max_rows_per_sheet=10))

    assert result["sheets"][0]["truncated"] is True
    assert described[0]["rows"] == 10 and [r["n"] for r in read(out)] == list(range(10))


def test_more_rows_than_one_batch_are_all_written(tmp_path) -> None:
    path = tmp_path / "many.csv"
    path.write_text("n,t\n" + "\n".join(f"{i},x{i}" for i in range(45_000)) + "\n", encoding="utf-8")

    _, described, out = tables_of(path, "many.csv", tmp_path)

    assert described[0]["rows"] == 45_000
    assert pq.read_metadata(out / "0.parquet").num_rows == 45_000
    assert read(out)[-1] == {"n": 44_999, "t": "x44999"}


@pytest.mark.parametrize(
    ("type_counts", "expected"),
    [
        ({"integer": 3}, "integer"),
        ({"integer": 2, "number": 1}, "number"),
        ({"number": 2, "error": 1}, "number"),
        ({"date": 2, "datetime": 1}, "datetime"),
        ({"text": 4}, "text"),
        ({"text": 3, "number": 1}, "text"),
        ({"error": 2}, "text"),
        ({}, "text"),
    ],
)
def test_column_kind_is_native_only_when_every_value_shares_it(type_counts, expected) -> None:
    assert column_kind({"type_counts": type_counts}) == expected
